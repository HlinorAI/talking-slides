"""Рендер слайдов в PDF (reportlab) и PPTX (python-pptx).

Оба формата оформлены темой из themes.py: градиентный фон с декором
(backdrop.py), цвет текста/акцентов из палитры темы.
"""

import os

from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas
from reportlab.lib.utils import simpleSplit
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

import themes as themes_mod
from backdrop import make_backdrop
from slidebuilder import format_ts

# TTF для встраивания в PDF: первый существующий путь из кандидатов.
_FONT_CANDIDATES = {
    "SlideFont": [
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "C:/Windows/Fonts/arial.ttf",
    ],
    "SlideFont-Bold": [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
        "C:/Windows/Fonts/arialbd.ttf",
    ],
}

PAGE_W, PAGE_H = landscape(A4)
MARGIN = 56


def _ensure_fonts():
    if "SlideFont" in pdfmetrics.getRegisteredFontNames():
        return
    for name, candidates in _FONT_CANDIDATES.items():
        path = next((p for p in candidates if os.path.exists(p)), None)
        if path is None:
            raise FileNotFoundError(
                f"{name}: не найден TTF-шрифт (Arial, DejaVu или Liberation). "
                "На Linux: sudo apt install fonts-dejavu-core")
        pdfmetrics.registerFont(TTFont(name, path))


def _rgb(c):
    return (c[0] / 255.0, c[1] / 255.0, c[2] / 255.0)


def _hex(c):
    return "%02X%02X%02X" % c


def _slides_word_ru(n):
    """Русская форма: 1 содержательный слайд / 2–4 слайда / 5+ слайдов."""
    if n % 10 == 1 and n % 100 != 11:
        return "содержательный слайд"
    if n % 10 in (2, 3, 4) and n % 100 not in (12, 13, 14):
        return "содержательных слайда"
    return "содержательных слайдов"


def _subtitle(slides, content_lang="en"):
    n = len(slides)
    if content_lang == "ru":
        sub = f"{n} {_slides_word_ru(n)}"
    else:
        sub = f"{n} content slide" + ("" if n == 1 else "s")
    if slides and slides[0].get("start") is not None:
        sub += f" · {format_ts(slides[0]['start'])}–{format_ts(slides[-1]['end'])}"
    return sub


def render_pdf(slides, path, title, source=None, theme_name=themes_mod.DEFAULT_THEME,
               content_lang="en"):
    _ensure_fonts()
    theme = themes_mod.resolve(theme_name)
    accent, text_c, muted = _rgb(theme["accent"]), _rgb(theme["text"]), _rgb(theme["muted"])
    bg = make_backdrop(theme, w=1684, h=1190, seed=theme_name)

    c = canvas.Canvas(path, pagesize=(PAGE_W, PAGE_H))
    c.setTitle(title)

    def draw_bg():
        c.drawImage(bg, 0, 0, PAGE_W, PAGE_H)

    # титульный слайд
    draw_bg()
    c.setFillColorRGB(*accent)
    c.rect(PAGE_W / 2 - 60, PAGE_H / 2 + 92, 120, 7, stroke=0, fill=1)
    c.setFillColorRGB(*text_c)
    c.setFont("SlideFont-Bold", 40)
    lines = simpleSplit(title, "SlideFont-Bold", 40, PAGE_W - 2 * MARGIN)
    y = PAGE_H / 2 + (len(lines) - 1) * 24 + 34
    for line in lines:
        c.drawCentredString(PAGE_W / 2, y, line)
        y -= 48
    c.setFont("SlideFont", 16)
    c.setFillColorRGB(*muted)
    c.drawCentredString(PAGE_W / 2, PAGE_H / 2 - len(lines) * 28 - 24,
                        _subtitle(slides, content_lang))
    if source:
        c.setFillAlpha(0.75)
        c.setFont("SlideFont", 11)
        c.drawCentredString(PAGE_W / 2, 34, os.path.basename(source))
        c.setFillAlpha(1)
    c.showPage()

    for i, slide in enumerate(slides, 1):
        draw_bg()

        # крупный полупрозрачный номер слайда
        c.setFillAlpha(0.35)
        c.setFillColorRGB(*text_c)
        c.setFont("SlideFont-Bold", 30)
        c.drawRightString(PAGE_W - MARGIN, PAGE_H - MARGIN - 6, f"{i:02d}")
        c.setFillAlpha(1)

        # акцентная плашка над заголовком
        c.setFillColorRGB(*accent)
        c.rect(MARGIN, PAGE_H - MARGIN + 18, 44, 5, stroke=0, fill=1)

        c.setFillColorRGB(*text_c)
        c.setFont("SlideFont-Bold", 23)
        tlines = simpleSplit(slide["title"], "SlideFont-Bold", 23, PAGE_W - 2 * MARGIN)
        y = PAGE_H - MARGIN - 8
        for line in tlines[:3]:
            c.drawString(MARGIN, y, line)
            y -= 30
        y -= 16

        for b in slide["bullets"]:
            blines = simpleSplit(b, "SlideFont", 17, PAGE_W - 2 * MARGIN - 24)
            if y - len(blines) * 26 < 64:
                break
            c.setFillColorRGB(*accent)
            c.circle(MARGIN + 4, y + 6, 3.5, stroke=0, fill=1)
            c.setFillColorRGB(*text_c)
            for line in blines:
                c.setFont("SlideFont", 17)
                c.drawString(MARGIN + 20, y, line)
                y -= 26
            y -= 12

        # футер
        c.setFillAlpha(0.7)
        c.setFillColorRGB(*muted)
        c.setFont("SlideFont", 9)
        ts = (
            f"{format_ts(slide['start'])}–{format_ts(slide['end'])}"
            if slide.get("start") is not None
            else ""
        )
        c.drawString(MARGIN, 30, ts)
        c.drawRightString(PAGE_W - MARGIN, 30, f"{i}/{len(slides)}")
        if source:
            c.drawCentredString(PAGE_W / 2, 30, os.path.basename(source))
        c.setFillAlpha(1)
        c.showPage()

    c.save()


def _html_escape(s):
    import html
    return html.escape(s, quote=True)


def render_html(slides, path, title, source=None, theme_name=themes_mod.DEFAULT_THEME,
                audio_fragments=None, lang="en", content_lang="en"):
    """Офлайн-плеер презентации: один HTML + папка media/ с фрагментами.
    Тема оформления, навигация по главам, кнопка ▶ у каждого слайда.
    audio_fragments: пути фрагментов по слайдам (None — без аудио);
    lang — язык документа (определённый язык записи)."""
    import html as html_mod

    theme = themes_mod.resolve(theme_name)
    accent = _hex(theme["accent"])
    bg_top, bg_bottom = _hex(theme["bg_stops"][0][1]), _hex(theme["bg_stops"][-1][1])
    text_c = _hex(theme["text"])
    muted = _hex(theme["muted"])

    lang_attr = (lang or "ru").lower()
    nav = []
    body = []
    for i, slide in enumerate(slides, 1):
        ts = (
            f"{format_ts(slide['start'])}–{format_ts(slide['end'])}"
            if slide.get("start") is not None
            else ""
        )
        anchor = f"slide-{i}"
        nav_time = format_ts(slide["start"]) if slide.get("start") is not None else str(i)
        nav.append(
            f'<a href="#{anchor}"><span class="t">{_html_escape(nav_time)}</span>'
            f'{_html_escape(slide["title"][:70])}</a>'
        )
        bullets = "\n".join(f"<li>{_html_escape(b)}</li>" for b in slide["bullets"])
        has_own = bool(audio_fragments and i <= len(audio_fragments)
                       and audio_fragments[i - 1])
        audio = (
            f'<audio controls preload="none" src="media/slide_{i:02d}.m4a"></audio>'
            if has_own
            else ""
        )
        src = (
            f'<div class="src">{_html_escape(os.path.basename(source))}</div>'
            if source
            else ""
        )
        body.append(f"""
<article id="{anchor}">
  <div class="num">{i:02d}</div>
  <h2>{_html_escape(slide["title"])}</h2>
  <ul>{bullets}</ul>
  <footer><span class="ts">{_html_escape(ts)}</span>{audio}</footer>
  {src}
</article>""")

    nav_html = "\n".join(nav)
    doc = f"""<!DOCTYPE html>
<html lang="{lang_attr}"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html_mod.escape(title)}</title>
<style>
  :root {{ --accent: #{accent}; --text: #{text_c}; --muted: #{muted};
          --bg1: #{bg_top}; --bg2: #{bg_bottom}; }}
  * {{ box-sizing: border-box; }}
  body {{ margin: 0; font: 16px/1.55 -apple-system, 'Helvetica Neue', Arial, sans-serif;
         color: var(--text);
         background: linear-gradient(160deg, var(--bg1), var(--bg2)) fixed; }}
  .wrap {{ max-width: 880px; margin: 0 auto; padding: 32px 20px 80px; }}
  h1 {{ font-size: 34px; margin: 24px 0 6px; }}
  .sub {{ color: var(--muted); margin-bottom: 28px; }}
  nav {{ position: sticky; top: 0; backdrop-filter: blur(14px);
        background: rgba(0,0,0,.25); border-radius: 14px; padding: 12px 16px;
        margin-bottom: 36px; max-height: 40vh; overflow: auto; }}
  nav a {{ display: flex; gap: 12px; color: var(--text); text-decoration: none;
          padding: 4px 0; font-size: 14px; opacity: .85; }}
  nav a:hover {{ color: var(--accent); opacity: 1; }}
  nav .t {{ color: var(--accent); font-variant-numeric: tabular-nums;
           min-width: 44px; }}
  article {{ background: rgba(0,0,0,.22); border-radius: 18px;
            padding: 28px 32px; margin-bottom: 26px; position: relative; }}
  article .num {{ position: absolute; top: 22px; right: 28px; font-size: 26px;
                 font-weight: 700; opacity: .3; }}
  article h2 {{ margin: 0 40px 14px 0; font-size: 22px; }}
  article ul {{ margin: 0; padding-left: 22px; }}
  article li {{ margin: 8px 0; }}
  footer {{ display: flex; align-items: center; gap: 16px; margin-top: 18px; }}
  .ts {{ color: var(--muted); font-size: 13px; font-variant-numeric: tabular-nums; }}
  audio {{ height: 34px; flex: 1; max-width: 420px; }}
  audio::-webkit-media-controls-panel {{ background: var(--accent); }}
  .src {{ color: var(--muted); font-size: 12px; margin-top: 10px; }}
</style></head>
<body><div class="wrap">
<h1>{html_mod.escape(title)}</h1>
<div class="sub">{_subtitle(slides, content_lang)}</div>
<nav>{nav_html}</nav>
{''.join(body)}
</div></body></html>"""
    with open(path, "w", encoding="utf-8") as f:
        f.write(doc)


def render_pptx(slides, path, title, source=None, theme_name=themes_mod.DEFAULT_THEME,
                audio_fragments=None, content_lang="en"):
    """audio_fragments: список путей к аудиофрагментам по слайдам (или None).
    Для каждого слайда с фрагментом в правый нижний угол встраивается
    аудиообъект — слайд проигрывает свой кусок записи прямо в PowerPoint."""
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN

    theme = themes_mod.resolve(theme_name)
    accent_hex, text_hex, muted_hex = (
        RGBColor(*theme["accent"]),
        RGBColor(*theme["text"]),
        RGBColor(*theme["muted"]),
    )
    bg = make_backdrop(theme, w=1920, h=1080, seed=theme_name)

    prs = Presentation()
    # автор из шаблона python-pptx («Steve Canny») не должен попадать
    # в метаданные чужих презентаций
    prs.core_properties.author = "talking-slides"
    prs.core_properties.last_modified_by = "talking-slides"
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def add_bg(slide):
        slide.shapes.add_picture(bg, 0, 0, width=prs.slide_width, height=prs.slide_height)

    def add_rect(slide, x, y, w, h, color):
        from pptx.util import Inches as In

        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, In(x), In(y), In(w), In(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*color)
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def add_text(slide, x, y, w, h, runs, size, bold=False, color=text_hex,
                 align=PP_ALIGN.LEFT, line_spacing=1.0):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = align
        if isinstance(runs, str):
            runs = [(runs, color, bold)]
        for text, rcolor, rbold in runs:
            run = p.add_run()
            run.text = text
            run.font.size = Pt(size)
            run.font.bold = rbold
            run.font.color.rgb = rcolor
        if line_spacing != 1.0:
            p.line_spacing = line_spacing
        return box

    # титульный слайд
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_rect(s, 6.17, 2.35, 1.0, 0.07, theme["accent"])
    add_text(s, 0.9, 2.6, 11.5, 1.9, title, 44, bold=True, align=PP_ALIGN.CENTER,
             line_spacing=1.05)
    add_text(s, 0.9, 4.45, 11.5, 0.6, _subtitle(slides, content_lang), 18,
             color=muted_hex,
             align=PP_ALIGN.CENTER)
    if source:
        add_text(s, 0.9, 6.75, 11.5, 0.4, os.path.basename(source), 12,
                 color=muted_hex, align=PP_ALIGN.CENTER)

    for i, slide in enumerate(slides, 1):
        s = prs.slides.add_slide(blank)
        add_bg(s)
        add_text(s, 11.4, 0.28, 1.5, 0.6, f"{i:02d}", 26, bold=True, color=muted_hex,
                 align=PP_ALIGN.RIGHT)
        add_rect(s, 0.7, 0.62, 0.38, 0.045, theme["accent"])
        add_text(s, 0.7, 0.78, 11.9, 1.3, slide["title"], 28, bold=True,
                 line_spacing=1.05)
        body = s.shapes.add_textbox(Inches(0.7), Inches(2.05), Inches(11.9), Inches(4.6))
        tf = body.text_frame
        tf.word_wrap = True
        first = True
        for b in slide["bullets"]:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            marker = p.add_run()
            marker.text = "•  "
            marker.font.size = Pt(20)
            marker.font.bold = True
            marker.font.color.rgb = accent_hex
            run = p.add_run()
            run.text = b
            run.font.size = Pt(20)
            run.font.color.rgb = text_hex
            p.space_after = Pt(14)

        ts = (
            f"{format_ts(slide['start'])}–{format_ts(slide['end'])}"
            if slide.get("start") is not None
            else ""
        )
        add_text(s, 0.7, 6.92, 3, 0.35, ts, 11, color=muted_hex)
        add_text(s, 11.2, 6.92, 1.4, 0.35, f"{i}/{len(slides)}", 11, color=muted_hex,
                 align=PP_ALIGN.RIGHT)
        if audio_fragments and i <= len(audio_fragments) and audio_fragments[i - 1]:
            s.shapes.add_movie(
                audio_fragments[i - 1],
                Inches(12.7), Inches(6.85), Inches(0.42), Inches(0.42),
                mime_type="audio/mp4",
            )
        if source:
            add_text(s, 4.17, 6.92, 5, 0.35, os.path.basename(source), 11,
                     color=muted_hex, align=PP_ALIGN.CENTER)

    prs.save(path)
